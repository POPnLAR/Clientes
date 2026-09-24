"""
Prepara la entrega semanal de leads de CLÍNICAS para el vendedor: selecciona los que recibieron la
secuencia automática completa y no respondieron, los limpia y los deja en un Excel listo para enviar.

Funciona en DOS pasos, para no equivocarse:

    python entregar_semana.py --cantidad 30                      # 1) VISTA PREVIA: no cambia nada
    python entregar_semana.py --cantidad 30 --confirmar --vendedor "Nombre"   # 2) registra la entrega

La vista previa muestra qué se entregaría y por qué se excluyó el resto, y genera un Excel de prueba.
Con --confirmar se genera el Excel definitivo y se registran los leads en asignaciones_vendedor.csv, para
que los workers automáticos dejen de escribirles.

Por seguridad se consulta al agente (VPS) qué leads ya respondieron; sin esa consulta no se entrega
nada, salvo que se indique --sin-agente (solo para pruebas). Las credenciales van en .env.entrega:
    AGENT_SERVICE_URL=https://gestionvitalpro.cl/agente-api
    AGENT_SERVICE_TOKEN=...
"""
import argparse
import datetime as dt
import os
import re
import subprocess
import sys

import pandas as pd

import agent_client
import asignaciones
import filtros_leads
from evo_client import es_movil_chileno, normalizar_telefono_chile

ARCHIVO_LEADS = "prospeccion_gestionvital_pro.csv"
CARPETA = "entregas"
LINEA = "clinicas"
RUBROS_PREFERIDOS = {"Clinica Estetica", "Medicina Estetica"}
ESTADOS_VENDEDOR = "Sin contactar,Contactado - sin respuesta,Interesado,Demo agendada,No interesado,Número incorrecto"

# El nombre debe sugerir estética/salud; si no, va a revisión manual en vez de al vendedor.
_RE_RUBRO_OK = re.compile(
    r"estetic|belleza|beauty|\bspa\b|cosmet|depila|laser|\bpiel\b|skin|facial|corporal|masaje|micropigment|"
    r"cejas|pestan|lash|clinic|medic|dermat|salon|nail|hands|feet|wellness|bella|studio|estudio|centro|"
    r"\bdra?\.|doctor|kine|glow|brow|hair|pelu|manicur|pedicur|podolog|estilo|salud|bienestar|rejuven|"
    r"lifting|botox|relleno|slim|body|cuerpo|unas|uñas|maquill|makeup|spa",
    re.IGNORECASE,
)
# Negocios de otro rubro que suelen colarse en las búsquedas (p. ej. "Clínica Veterinaria").
_RE_OTRO_RUBRO = re.compile(
    r"veterin|mascota|optic|farmac|inmobil|hotel|ferreter|restaurant|panader|supermerc|abogad|contab|mecanic",
    re.IGNORECASE,
)
_RE_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _sin_acentos(t):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn").lower()


def limpiar_nombre(nombre):
    """Espacios normalizados, sin colas de SEO, largo acotado; TODO EN MAYÚSCULAS pasa a formato título."""
    n = re.sub(r"\s+", " ", str(nombre or "")).strip()
    n = n.split("|")[0].strip() or n          # "Marca | servicio | servicio ..." -> "Marca"
    if len(n) > 60:
        n = n[:60].rsplit(" ", 1)[0].rstrip(" ,-—") + "…"
    return n.title() if len(n) > 4 and n.isupper() else n


def formatear_telefono(tel):
    d = "".join(filter(str.isdigit, normalizar_telefono_chile(tel)))
    return f"+56 9 {d[3:7]} {d[7:]}" if es_movil_chileno(d) else str(tel)


def rubro_de(ministerio):
    return str(ministerio or "").split(" - ", 1)[-1].strip() or "Estética"


def cargar_env(path):
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as f:
        for linea in f:
            linea = linea.strip()
            if linea and not linea.startswith("#") and "=" in linea:
                k, v = linea.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def seleccionar(df, respondieron, asignados, hoy, min_dias=3):
    """
    Devuelve (candidatos_df, motivos, revisar_df). Un lead califica si recibió la secuencia completa
    (Estado Finalizado), no tiene un resultado registrado, tiene celular y no está en cadenas,
    respuestas ni asignaciones. `motivos` cuenta por qué se excluyó el resto.
    """
    motivos = {}

    def excluir(mask, motivo):
        nonlocal df
        n = int(mask.sum())
        if n:
            motivos[motivo] = motivos.get(motivo, 0) + n
        df = df[~mask]

    df = df.copy().fillna("")
    df["_clave"] = df["Telefono"].map(asignaciones.clave_telefono)
    df["_tel"] = df["Telefono"].map(normalizar_telefono_chile)
    ultimo = pd.to_datetime(df["Fecha_Contacto"], format="%d/%m/%Y %H:%M", errors="coerce")
    df["_ultimo"] = ultimo

    excluir(df["Estado"] != "Finalizado", "no terminó la secuencia (Nuevo/Contactado/Rechazado/Error/Agendado)")
    excluir(df["Resultado"].str.strip() != "", "ya tiene un resultado registrado (no interesado, agenda, otro rubro...)")
    excluir(df["Notas"].str.contains("Respondi", case=False), "respondió por WhatsApp")
    excluir(~df["_tel"].map(es_movil_chileno), "sin celular")
    excluir(df["Evento"].map(lambda x: filtros_leads.es_cadena(x, LINEA)), "cadena o institución")
    excluir(df["_clave"].isin(asignados), "ya entregado antes al vendedor")
    excluir(df["_clave"].isin(respondieron), "ya respondió al agente")
    excluir(df["_ultimo"].isna() | ((hoy - df["_ultimo"]).dt.days < min_dias), f"último mensaje hace menos de {min_dias} días")

    excluir(df["Evento"].map(lambda x: bool(_RE_OTRO_RUBRO.search(_sin_acentos(x)))), "otro rubro (veterinaria, óptica, farmacia...)")

    dudoso = ~df["Evento"].map(lambda x: bool(_RE_RUBRO_OK.search(_sin_acentos(x))))
    revisar = df[dudoso].copy()
    if len(revisar):
        motivos["nombre que no parece de estética (para revisar)"] = len(revisar)
    df = df[~dudoso]

    df["_email"] = df.apply(
        lambda r: r["Email"].strip().lower()
        if _RE_EMAIL.match(r["Email"].strip()) and "invalido" not in r["Notas"].lower() else "",
        axis=1,
    )
    df["_pref"] = df["Ministerio"].map(rubro_de).isin(RUBROS_PREFERIDOS)
    df = df.sort_values(["_email", "_pref", "_ultimo"], key=lambda c: c.map(bool) if c.name == "_email" else c,
                        ascending=[False, False, False])
    return df, motivos, revisar


def construir_filas(sel):
    filas = []
    for i, (_, r) in enumerate(sel.iterrows(), start=1):
        enviados = int(float(r.get("Dia_Secuencia") or 0))
        ultimo = r["_ultimo"].strftime("%d/%m/%Y")
        digitos = "".join(filter(str.isdigit, r["_tel"]))
        filas.append({
            "N°": i,
            "Prioridad": "A" if r["_email"] else "B",
            "Negocio": limpiar_nombre(r["Evento"]),
            "Rubro": rubro_de(r["Ministerio"]),
            "Zona de búsqueda": str(r["Ubicacion"]).strip(),
            "Teléfono": formatear_telefono(r["Telefono"]),
            "WhatsApp": f"https://wa.me/{digitos}",
            "Email": r["_email"],
            "Último contacto": ultimo,
            "Contexto": f"Recibió {enviados} mensajes automáticos por WhatsApp (último el {ultimo}); no ha respondido.",
            "Estado": "Sin contactar",
            "Fecha de contacto": "",
            "Comentarios": "",
            "_telefono_crudo": r["Telefono"],
        })
    return filas


def escribir_excel(filas, ruta, vendedor, hoy):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    columnas = [c for c in filas[0].keys() if not c.startswith("_")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(columnas)
    for f in filas:
        ws.append([f[c] for c in columnas])

    relleno = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    anchos = {"N°": 5, "Prioridad": 10, "Negocio": 34, "Rubro": 20, "Zona de búsqueda": 18, "Teléfono": 17, "WhatsApp": 14,
              "Email": 30, "Último contacto": 14, "Contexto": 52, "Estado": 24, "Fecha de contacto": 16, "Comentarios": 44}
    for i, nombre in enumerate(columnas, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = anchos.get(nombre, 16)
        for celda in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            celda[0].alignment = Alignment(vertical="top", wrap_text=nombre in ("Contexto", "Comentarios", "Negocio"))
    col_wa = columnas.index("WhatsApp") + 1
    for fila in range(2, len(filas) + 2):
        c = ws.cell(fila, col_wa)
        c.hyperlink = c.value
        c.value = "Abrir chat"
        c.font = Font(color="0563C1", underline="single")
    dv = DataValidation(type="list", formula1=f'"{ESTADOS_VENDEDOR}"', allow_blank=True)
    ws.add_data_validation(dv)
    col_estado = columnas.index("Estado") + 1
    dv.add(f"{ws.cell(2, col_estado).coordinate}:{ws.cell(len(filas) + 1, col_estado).coordinate}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    info = wb.create_sheet("Cómo usar")
    texto = [
        f"Entrega del {hoy.strftime('%d/%m/%Y')} — vendedor: {vendedor}",
        "",
        "Qué es esta lista:",
        "  Negocios de estética de la Región Metropolitana que ya recibieron mensajes automáticos de GestiónVital Pro por WhatsApp y NO han respondido.",
        "",
        "Prioridad:",
        "  A = tiene email además de teléfono (dos vías de contacto).   B = solo teléfono/WhatsApp.",
        "",
        "Cómo trabajarla:",
        "  1. Contacta a cada negocio (el link de la columna WhatsApp abre el chat).",
        "  2. Actualiza la columna Estado (lista desplegable), la fecha de contacto y los comentarios.",
        "  3. Si el negocio dice que no le interesa, márcalo 'No interesado' y no insistas.",
        "  4. Devuelve el archivo actualizado al final de la semana.",
        "",
        "No contactes a negocios que no estén en tu lista: pueden estar siendo atendidos por otra vía.",
        "",
        "Ojo con la 'Zona de búsqueda': es la comuna en la que se buscó el negocio, NO su dirección exacta (Google a veces",
        "devuelve negocios de zonas cercanas). Confirma la ubicación al hablar con ellos.",
    ]
    for t in texto:
        info.append([t])
    info.column_dimensions["A"].width = 130
    info["A1"].font = Font(bold=True, size=13)
    wb.save(ruta)


def auditar(df):
    """Leads de TODA la base (cualquier estado) cuyo nombre es de otro rubro o no sugiere estética."""
    nombres = df["Evento"].map(_sin_acentos)
    otro = nombres.map(lambda x: bool(_RE_OTRO_RUBRO.search(x)))
    dudoso = ~nombres.map(lambda x: bool(_RE_RUBRO_OK.search(x)))
    out = df[otro | dudoso].copy()
    out["_motivo"] = ["OTRO RUBRO" if o else "nombre sin pista de estética" for o in otro[out.index]]
    return out


def base_al_dia():
    """Avisa si la copia local está atrasada respecto de GitHub (los workers actualizan el CSV cada hora)."""
    try:
        subprocess.run(["git", "fetch", "-q", "origin"], capture_output=True, timeout=30)
        n = subprocess.run(["git", "rev-list", "--count", "HEAD..origin/main"], capture_output=True, text=True, timeout=30)
        return int(n.stdout.strip() or 0) == 0
    except Exception:
        return True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cantidad", type=int, default=30, help="cuántos leads entregar (por defecto 30)")
    ap.add_argument("--vendedor", default="Vendedor", help="nombre del vendedor (queda registrado)")
    ap.add_argument("--min-dias", type=int, default=3, help="días mínimos desde el último mensaje automático")
    ap.add_argument("--confirmar", action="store_true", help="registra la entrega (sin esto es solo vista previa)")
    ap.add_argument("--sin-agente", action="store_true", help="NO verificar respuestas en el agente (solo pruebas)")
    ap.add_argument("--env", default=".env.entrega")
    ap.add_argument("--auditar", action="store_true",
                    help="lista los leads de otro rubro o de nombre dudoso en TODA la base (no entrega nada)")
    args = ap.parse_args()

    if args.auditar:
        base = pd.read_csv(ARCHIVO_LEADS, dtype=str).fillna("")
        a = auditar(base)
        print(f"Leads de clínicas de otro rubro o de nombre dudoso: {len(a)} de {len(base)}\n")
        print("⚠️  Los que están en Nuevo o Contactado SIGUEN recibiendo mensajes automáticos.")
        print("    Para descartarlos: dashboard > Editor de Base > Estado = Rechazado.\n")
        print(f"{'Id':>4} {'Estado':11} {'Negocio':46} {'Zona':15} Motivo")
        orden = {"Nuevo": 0, "Contactado": 1}
        for _, r in a.sort_values("Estado", key=lambda c: c.map(lambda x: orden.get(x, 2))).iterrows():
            aviso = "  <-- ACTIVO" if r["Estado"] in orden else ""
            print(f"{r['Id']:>4} {r['Estado']:11} {r['Evento'][:45]:46} {r['Ubicacion'][:14]:15} {r['_motivo']}{aviso}")
        return

    cargar_env(args.env)
    if not base_al_dia():
        print("⚠️  Tu copia local está ATRASADA respecto de GitHub. Ejecuta `git pull` y vuelve a correr esto: "
              "los workers actualizan los leads cada hora.\n")

    if args.sin_agente:
        respondieron = set()
        print("⚠️  Sin verificar respuestas en el agente (--sin-agente): NO uses este resultado para entregar.\n")
    else:
        r = agent_client.obtener_respuestas()
        if r["estado"] != "ok":
            sys.exit(f"No se pudo consultar al agente (estado: {r['estado']}). Revisa AGENT_SERVICE_URL y "
                     f"AGENT_SERVICE_TOKEN en {args.env}. Sin saber quién ya respondió no se entrega nada.")
        respondieron = r["humanos"] | r["bots"]
        print(f"Agente consultado: {len(r['humanos'])} contactos que respondieron y {len(r['bots'])} con respuesta de bot.\n")

    df = pd.read_csv(ARCHIVO_LEADS, dtype=str)
    hoy = pd.Timestamp(dt.datetime.now())
    asignados = asignaciones.cargar_asignados()
    sel, motivos, revisar = seleccionar(df, respondieron, asignados, hoy, args.min_dias)

    print(f"Base de clínicas: {len(df)} leads")
    for m, n in sorted(motivos.items(), key=lambda x: -x[1]):
        print(f"  - excluidos {n:3}  {m}")
    print(f"  = disponibles para entregar: {len(sel)}")
    if len(revisar):
        print("\nNombres que NO parecen de estética (revísalos, no van al vendedor):")
        for _, r in revisar.iterrows():
            print(f"    · {r['Evento'][:45]:45} {r['Ubicacion']}  ({r['Telefono']})")

    elegidos = sel.head(args.cantidad)
    if elegidos.empty:
        sys.exit("\nNo hay leads disponibles para entregar con estas reglas.")
    filas = construir_filas(elegidos)
    print(f"\nEntrega propuesta: {len(filas)} leads  (A = con email: {sum(f['Prioridad'] == 'A' for f in filas)}, B: {sum(f['Prioridad'] == 'B' for f in filas)})")
    for f in filas[:8]:
        print(f"    {f['N°']:2} [{f['Prioridad']}] {f['Negocio'][:36]:36} {f['Zona de búsqueda'][:14]:14} {f['Teléfono']}  {f['Rubro']}")
    if len(filas) > 8:
        print(f"    ... y {len(filas) - 8} más")

    os.makedirs(CARPETA, exist_ok=True)
    hoy_txt = hoy.strftime("%Y-%m-%d")
    if not args.confirmar:
        ruta = os.path.join(CARPETA, f"VISTA_PREVIA_{hoy_txt}.xlsx")
        escribir_excel(filas, ruta, args.vendedor, hoy)
        print(f"\nVISTA PREVIA guardada en {ruta} (no se registró ninguna entrega).")
        print("Si está bien, repite el comando agregando:  --confirmar --vendedor \"Nombre\"")
        return

    nombre_arch = re.sub(r"[^A-Za-z0-9_-]+", "_", _sin_acentos(args.vendedor)).strip("_") or "vendedor"
    ruta = os.path.join(CARPETA, f"entrega_{hoy_txt}_{nombre_arch}.xlsx")
    escribir_excel(filas, ruta, args.vendedor, hoy)
    semana = hoy.strftime("%G-W%V")
    n = asignaciones.registrar([
        {"Telefono": f["_telefono_crudo"], "Linea": LINEA, "Nombre": f["Negocio"], "Asignado_A": args.vendedor,
         "Fecha_Asignacion": hoy_txt, "Semana": semana}
        for f in filas
    ])
    print(f"\n✅ Entrega registrada: {n} leads asignados a {args.vendedor}.")
    print(f"   Excel para enviar: {ruta}")
    print("\nÚLTIMO PASO (para que los workers dejen de escribirles a estos leads):")
    print("   git checkout main && git pull")
    print("   git add asignaciones_vendedor.csv")
    print(f"   git commit -m \"Asignacion semanal al vendedor ({semana})\"")
    print("   git push origin main")


if __name__ == "__main__":
    main()
